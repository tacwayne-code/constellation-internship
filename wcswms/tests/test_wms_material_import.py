"""Material import uses isolated ledgers and a FakePLC; no motion or stock posting."""
import base64
import copy
import io

import pytest
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient as AnonymousClient
from wms_test_client import TestClient, PASSWORD
from backend.wms_app import create_app
from backend.wms_store import WmsStore
from test_wms import CONFIG, FakePLC, request

HEADER = '物料编码,物料名称,型号,规格\n'

@pytest.fixture
def system(tmp_path):
    plc = FakePLC()
    app = create_app(config=copy.deepcopy(CONFIG), data_dir=tmp_path, reader=plc,
                     commands_path=tmp_path/'commands.db', automatic=False, key='qa-key')
    return app, TestClient(app), plc

def file(text, encoding='utf-8-sig'):
    return dict(filename='物料.csv', content=base64.b64encode(text.encode(encoding)).decode())

def excel(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    result = io.BytesIO(); book.save(result)
    return dict(filename='物料.xlsx', content=base64.b64encode(result.getvalue()).decode())

def preview(client, payload, prefix=''):
    result = client.post(prefix+'/api/materials/import/preview', json=payload)
    assert result.status_code == 200, result.text
    return result.json()

def commit(client, payload, checked, prefix='', **changes):
    return client.post(prefix+'/api/materials/import', json={**payload, 'request_id':'test-import-once',
        'expected_token':checked['token'], 'confirmed':True, **changes})

def test_preview_read_only_import_zero_stock_lookup_and_idempotent_restart(system):
    app, client, plc = system; store = app.state.store
    payload = file(HEADER+'0006204,轴承,6204,20 × 47 mm\nABC,传动组件,,\n')
    before = store.read(); checked = preview(client, payload)
    assert store.read() == before
    assert checked['counts'] == dict(NEW=2, FILL=0, UNCHANGED=0, ERROR=0)
    assert commit(client, payload, checked, confirmed=False).status_code == 409
    result = commit(client, payload, checked)
    assert result.status_code == 200, result.text
    after = store.read()
    assert commit(client, payload, checked).json() == result.json()
    assert store.read() == after
    assert after['stock'] == before['stock'] and after['tasks'] == before['tasks'] and plc.writes == []
    assert client.get('/api/materials/lookup?sku=0006204').json()['material']['material_name'] == '轴承'
    assert len(client.get('/api/materials').json()['materials']) == 2
    assert WmsStore(store.path, CONFIG).read()['materials']['0006204']['model'] == '6204'
    assert commit(client, file(HEADER+'NEW,改动,,\n'), checked).status_code == 409

def test_excel_template_text_codes_and_no_demo_material(system):
    _, client, _ = system
    response = client.get('/api/materials/template.xlsx')
    assert response.status_code == 200 and 'attachment' in response.headers['content-disposition']
    book = load_workbook(io.BytesIO(response.content))
    assert book.active['A2'].number_format == '@' and book.active['A2'].value is None
    assert book.active['A1'].value == '物料编码'
    checked = preview(client, excel([['物料编码','物料名称','型号','规格'],['000123','电机','M1','30mm']]))
    assert checked['rows'][0]['sku'] == '000123' and checked['valid']
    numeric = preview(client, excel([['物料编码','物料名称'],[123,'电机']]))
    assert not numeric['valid'] and '文本' in numeric['rows'][0]['errors'][0]
    missing = preview(client, excel([['物料编码','物料名称'],[None,'空编码']]))
    assert not missing['valid'] and missing['rows'][0]['row'] == 2

@pytest.mark.parametrize('encoding', ['utf-8-sig','gb18030'])
def test_csv_quoted_names_newlines_blank_lines_and_english_headers(system, encoding):
    _, client, _ = system
    payload = file('sku,material_name,model,specification\n\n001,"轴承,组件",M,"第一行\n第二行"\n', encoding)
    checked = preview(client, payload)
    assert checked['valid'] and checked['total'] == 1
    assert checked['rows'][0]['material_name'] == '轴承,组件'
    assert checked['rows'][0]['specification'] == '第一行\n第二行'

@pytest.mark.parametrize('text', [
    HEADER+'A,轴承,,\nA,重复,,\n',
    HEADER+'A,轴承,,\n,空编码,,\n',
    HEADER+'A,轴承,,\nB,'+'X'*121+',,\n',
    HEADER+'A,轴承,,\nB,额外,,,列\n',
])
def test_error_rows_prevent_entire_batch_write(system, text):
    app, client, plc = system
    payload = file(text); before = app.state.store.read(); checked = preview(client, payload)
    assert not checked['valid'] and checked['counts']['ERROR'] > 0
    assert commit(client, payload, checked).status_code == 409
    assert app.state.store.read() == before and plc.writes == []

def test_fill_stock_metadata_preserve_inventory_and_task_snapshots(system):
    app, client, plc = system; store = app.state.store
    store.verify_empty('L-01-01-1')
    store.stocktake(dict(request_id='seed-for-import',barcode='BOX',sku='CODE',material_name='轴承',
                         quantity=7,batch='OLD',location='L-01-01-1',note='核对'))
    store.create_task(request(kind='OUTBOUND',barcode='BOX',material_name='轴承') | dict(sku='CODE',quantity=7,batch='OLD'))
    before = store.read()
    payload = file(HEADER+'CODE,,6204,20mm\n')
    checked = preview(client, payload)
    assert checked['counts']['FILL'] == 1
    assert commit(client, payload, checked).status_code == 200
    after = store.read(); old, new = before['stock']['BOX'], after['stock']['BOX']
    assert new['material_name'] == '轴承' and new['model'] == '6204'
    for key in ('quantity','location','batch','status','reserved'):
        assert old[key] == new[key]
    assert after['tasks'] == before['tasks'] and plc.writes == []
    checked = preview(client, payload)
    assert checked['counts']['UNCHANGED'] == 1
    assert commit(client, payload, checked, request_id='unchanged-import').status_code == 200

def test_conflict_and_catalog_change_between_preview_and_commit(system):
    app, client, _ = system
    app.state.store.mutate(lambda state: state['materials'].update(CODE=dict(sku='CODE',material_name='旧名称',model='',specification='')))
    conflict = file(HEADER+'CODE,新名称,,\nNEW,新增,,\n')
    checked = preview(client, conflict)
    assert checked['counts']['ERROR'] == 1 and commit(client, conflict, checked).status_code == 409
    payload = file(HEADER+'NEW,新增,,\n'); checked = preview(client, payload)
    app.state.store.mutate(lambda state: state['materials'].update(NEW=dict(sku='NEW',material_name='其他',model='',specification='')))
    result = commit(client, payload, checked)
    assert result.status_code == 409 and '重新预览' in result.text

def test_unfinished_material_order_blocks_fill_until_cancelled(system):
    app, client, _ = system; store = app.state.store
    store.verify_empty('L-01-01-1')
    store.stocktake(dict(request_id='seed-empty-for-import',barcode='EMPTY',container_type='EMPTY_BIN',sku='',
                         quantity=0,batch='',location='L-01-01-1',note='核对'))
    order = store.orders.create(dict(request_id='order-for-import',kind='INBOUND',sku='CODE',material_name='轴承',
                                    model='',specification='',quantity=3,batch='',box_mode='NEW',box_barcode=''))
    payload = file(HEADER+'CODE,,6204,20mm\n')
    checked = preview(client,payload)
    assert checked['counts']['ERROR'] == 1 and '未结束' in checked['rows'][0]['errors'][0]
    assert commit(client,payload,checked).status_code == 409
    store.cancel(order['current_task_id'])
    checked = preview(client,payload)
    assert checked['counts']['FILL'] == 1 and commit(client,payload,checked).status_code == 200

def test_exact_batch_limit_and_size_rejection(system):
    app, client, _ = system
    payload = file(HEADER+'\n'.join(f'{index:06},轴承,,' for index in range(1000)))
    checked = preview(client,payload)
    assert checked['total'] == 1000 and checked['valid']
    assert commit(client,payload,checked).status_code == 200
    before = app.state.store.read()
    huge = dict(filename='large.csv',content=base64.b64encode(b'a'*(2*1024*1024+1)).decode())
    assert client.post('/api/materials/import/preview',json=huge).status_code == 409
    assert app.state.store.read() == before

@pytest.mark.parametrize('payload', [
    file('物料名称,型号\n轴承,M\n'), file('sku,物料编码\nA,B\n'), file(HEADER),
    file(HEADER+'\n'.join(f'{index},名称,,' for index in range(1001))),
    dict(filename='物料.xls',content='YQ=='), dict(filename='物料.csv',content='not-base64'),
    dict(filename='物料.xlsx',content='YQ=='),
    excel([['物料编码'],['=1+1']]),
])
def test_invalid_file_rejected_without_change(system, payload):
    app, client, _ = system; before = app.state.store.read()
    result = client.post('/api/materials/import/preview', json=payload)
    assert result.status_code == 409, result.text
    assert app.state.store.read() == before

def test_auth_permissions_csrf_and_device_scope(system):
    app, client, plc = system
    payload = file(HEADER+'LOCAL,当前设备,,\n')
    assert AnonymousClient(app).get('/api/materials').status_code == 401
    assert client.post('/api/materials/import/preview',json=payload,headers={'X-CSRF-Token':''}).status_code == 403
    for role in ('viewer','operator'):
        client.post('/api/users',json=dict(username=role,display_name=role,password=PASSWORD,role=role))
        other = AnonymousClient(app)
        login = other.post('/api/auth/login',json=dict(username=role,password=PASSWORD),headers={'X-WMS-Request':'1'})
        other.headers['X-CSRF-Token'] = login.json()['csrf_token']
        assert other.get('/api/materials').status_code == 200
        assert other.post('/api/materials/import/preview',json=payload).status_code == 403
        assert commit(other,payload,dict(token='a'*64)).status_code == 403
    device = client.post('/api/devices',json=dict(warehouse_id='main',name='导入隔离设备',host='192.168.1.151')).json()
    prefix = '/api/scopes/'+device['id']
    checked = preview(client,payload,prefix)
    assert commit(client,payload,checked,prefix).status_code == 200
    assert client.get('/api/materials').json()['materials'] == []
    assert client.get(prefix+'/api/materials').json()['materials'][0]['sku'] == 'LOCAL'
    assert plc.writes == []
